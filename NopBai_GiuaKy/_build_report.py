from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn

ROOT = Path(r"D:\LAB\PTUDD\NopBai_GiuaKy")
IMAGES = ROOT / "images"
OUT = ROOT / "Bao_cao_ket_qua_hai_ung_dung.docx"

doc = Document()
doc.core_properties.author = ""
doc.core_properties.comments = ""
doc.core_properties.title = "Báo cáo kết quả hai ứng dụng quản lý sinh viên"
section = doc.sections[0]
section.top_margin = Cm(2.0)
section.bottom_margin = Cm(1.8)
section.left_margin = Cm(2.3)
section.right_margin = Cm(2.3)

styles = doc.styles
for name in ["Normal", "Title", "Heading 1", "Heading 2"]:
    style = styles[name]
    style.font.name = "Arial"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].paragraph_format.space_after = Pt(5)
styles["Title"].font.size = Pt(19)
styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
title_properties = styles["Title"]._element.get_or_add_pPr()
for border in title_properties.findall(qn("w:pBdr")):
    title_properties.remove(border)
styles["Heading 1"].font.size = Pt(13)
styles["Heading 1"].paragraph_format.space_before = Pt(11)
styles["Heading 1"].paragraph_format.space_after = Pt(5)
styles["Heading 2"].font.size = Pt(11)
styles["Heading 2"].paragraph_format.space_before = Pt(7)
styles["Heading 2"].paragraph_format.space_after = Pt(3)

def para(text="", bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p

def image(filename, caption, width=15.8):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(IMAGES / filename), width=Cm(width))
    c = doc.add_paragraph(caption)
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_after = Pt(7)
    for run in c.runs:
        run.italic = True
        run.font.size = Pt(9)

def table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, value in enumerate(headers):
        t.rows[0].cells[i].text = value
        for run in t.rows[0].cells[i].paragraphs[0].runs:
            run.bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    doc.add_paragraph()

doc.add_paragraph("Báo cáo kết quả hai ứng dụng quản lý sinh viên", style="Title")
para("Bài kiểm tra thực hành giữa kỳ  Chủ đề 2 và Chủ đề 3")
para("Họ và tên: ........................................................    MSSV: ........................................................")
para("Lớp: ........................................................    Ngày trình bày: ........................................................")

doc.add_heading("Mục tiêu và phạm vi", level=1)
para("Báo cáo trình bày kết quả chạy của hai ứng dụng WinForms: Quản lý sinh viên khoa CNTT thuộc Chủ đề 2 và bài tập quản lý sinh viên, thao tác tập tin thuộc Chủ đề 3. Hai ứng dụng dùng đối tượng SinhVien để giữ dữ liệu trong bộ nhớ và ListView để hiển thị. Chủ đề 2 đọc và ghi tệp TXT; Chủ đề 3 hỗ trợ đọc, ghi JSON, XML và TXT.")

doc.add_heading("Cấu trúc chương trình", level=1)
table(["Thành phần", "Chủ đề 2", "Chủ đề 3"], [
    ["SinhVien", "Lưu MSSV, họ tên, ngày sinh, giới tính, lớp, liên hệ và hình.", "Lưu MSSV, họ tên lót, tên, lớp, CMND, liên hệ và môn học đăng ký."],
    ["Giao diện", "frmSinhVienCNTT nhận dữ liệu, hiển thị ListView và xử lý sự kiện nút bấm.", "frmQuanLySinhVien điều phối; frmTimKiem nhận điều kiện; frmNhapFile chọn định dạng và tệp."],
    ["Dữ liệu", "List<SinhVien> dsSV và tệp DSNV.txt.", "QuanLySinhVien giữ List<SinhVien>, đường dẫn và loại tệp."],
])
para("Luồng chung: người dùng nhập hoặc chọn dữ liệu trên form → đối tượng SinhVien → danh sách trong bộ nhớ → ListView. Khi lưu tệp, dữ liệu trong danh sách được tuần tự hóa; khi đọc tệp, danh sách được nạp lại rồi hiển thị.")

doc.add_heading("Ứng dụng Chủ đề 2  Quản lý sinh viên khoa CNTT", level=1)
image("chu_de_2_cntt_danh_sach.png", "Hình 1. Giao diện Chủ đề 2 sau khi đọc ba sinh viên mẫu từ DSNV.txt")
para("Form dùng MaskedTextBox cho MSSV và số điện thoại, DateTimePicker cho ngày sinh, ComboBox cho lớp, RadioButton cho giới tính, ListView cho danh sách. Cách chọn control giúp giới hạn kiểu nhập và cho phép quan sát nhiều bản ghi cùng lúc.")

doc.add_heading("Chức năng và kết quả Chủ đề 2", level=2)
table(["Thao tác", "Luồng xử lý trong mã nguồn"], [
    ["Mở ứng dụng", "frmSinhVienCNTT_Load gọi DocTuFile, nạp DSNV.txt cạnh ứng dụng vào dsSV rồi gọi LoadListView."],
    ["Lưu để thêm hoặc cập nhật", "GetSinhVienFromForm tạo SinhVien; TimSinhVien tìm theo MSSV. Không tìm thấy thì thêm, có thì cập nhật; sau đó tải lại ListView và đánh dấu isModified."],
    ["Xóa", "Menu chuột phải lấy MSSV của dòng chọn, XoaSinhVien xóa trong danh sách, rồi tải lại giao diện."],
    ["Tải lại", "Đọc lại DSNV.txt, thay dữ liệu trong bộ nhớ và bỏ các thay đổi chưa ghi."],
    ["Thoát", "Nếu isModified, hộp thoại hỏi lưu. Yes ghi DSNV.txt; No bỏ thay đổi; Cancel giữ form mở."],
])
image("chu_de_2_cntt_chon_sinh_vien.png", "Hình 2. Chọn một dòng sẽ đưa thông tin sinh viên lên các ô nhập")
image("chu_de_2_cntt_cap_nhat.png", "Hình 3. Sau khi bấm Lưu, tên được cập nhật ngay trên ListView; bước ghi tệp diễn ra khi thoát và chọn Yes")
para("MSSV là khóa nhận diện bản ghi trong TimSinhVien và XoaSinhVien. Vòng lặp xóa đi từ cuối danh sách để các chỉ số chưa xét không dịch chuyển. Dữ liệu được lưu theo các cột phân cách bằng tab, theo cùng thứ tự mà DocTuFile đọc.")

doc.add_heading("Ứng dụng Chủ đề 3  Quản lý sinh viên và tập tin", level=1)
image("chu_de_3_danh_sach.png", "Hình 4. Giao diện Chủ đề 3 sau khi nạp students.json; MSSV kế tiếp của CTK44 được gợi ý là 2010003", width=14.0)
para("Form chính kiểm tra đầu vào ở KiemTraNhap, tạo đối tượng ở LayThongTinForm và chuyển thao tác dữ liệu cho QuanLySinhVien. Lớp quản lý xử lý thêm, cập nhật, xóa, tìm kiếm, sinh MSSV và đọc/ghi tệp. Sau thao tác, form gọi HienThiDanhSach để ListView khớp với danh sách trong bộ nhớ.")

doc.add_heading("Chức năng và kết quả Chủ đề 3", level=2)
table(["Thao tác", "Luồng xử lý trong mã nguồn"], [
    ["Thêm", "Kiểm tra MSSV, tên, CMND, số điện thoại, địa chỉ, môn học; từ chối MSSV trùng; thêm vào danh sách và ghi tệp."],
    ["Cập nhật", "Tìm đúng MSSV, thay bản ghi trong danh sách và ghi tệp."],
    ["Xóa", "Lấy MSSV từ dòng được chọn hoặc tích, xóa trong danh sách và ghi tệp."],
    ["Tìm kiếm", "frmTimKiem nhận MSSV, tên, lớp; QuanLySinhVien.TimKiem lọc từng sinh viên; form hiển thị kết quả."],
    ["Nhập từ tệp", "frmNhapFile chọn JSON/XML/TXT và đường dẫn; DocTuFile đọc rồi làm mới ListView."],
])
image("chu_de_3_ket_qua_tim_kiem.png", "Hình 5. Tìm theo tên Nga trả về một sinh viên có MSSV 2010002")
image("chu_de_3_nhap_file.png", "Hình 6. Hộp thoại chọn định dạng và tệp nhập dữ liệu")
para("DuongDanFile và LoaiFile cho biết tệp nào sẽ được cập nhật sau khi thêm, sửa hoặc xóa. JSON dùng Newtonsoft.Json, XML dùng XmlSerializer, TXT dùng các trường phân cách bằng dấu *. Tệp mặc định students.json được tìm cạnh tệp chạy để ứng dụng không phụ thuộc thư mục làm việc hiện tại.")

doc.add_heading("Giải thích quyết định thiết kế", level=1)
design_first = para("Hai bài dùng List<SinhVien> vì số bản ghi của bài thực hành nhỏ, việc duyệt, thêm và xóa bằng vòng lặp dễ theo dõi. MSSV dùng để phân biệt sinh viên; họ tên có thể trùng. ListView chỉ là phần hiển thị, không phải nguồn dữ liệu chính. MaskedTextBox hỗ trợ nhập theo khuôn dạng, còn kiểm tra trong sự kiện nút bấm quyết định dữ liệu có được xử lý hay không.")
design_first.paragraph_format.keep_together = True
para("Chủ đề 2 lưu khi đóng để người dùng quyết định giữ hay bỏ thay đổi. Chủ đề 3 ghi tệp ngay sau các thao tác thành công để dữ liệu trên đĩa được cập nhật tức thời. Hai cách này có ưu và nhược điểm khác nhau: lưu khi đóng cho phép hủy thay đổi chưa lưu; ghi ngay giảm nguy cơ quên lưu nhưng cần kiểm soát đường dẫn tệp đang mở.")

doc.add_heading("Tình huống kiểm tra và hướng xử lý", level=1)
table(["Tình huống", "Dự đoán và vị trí xử lý"], [
    ["Nhập MSSV trùng ở Chủ đề 3", "Them trả false; giao diện báo trùng mã. Kiểm tra điều kiện trong QuanLySinhVien.Them."],
    ["Tìm theo tên không phân biệt chữ hoa", "TimKiem chuyển chuỗi về chữ thường trước khi so sánh."],
    ["Đổi lớp khi tạo sinh viên Chủ đề 3", "cboLop_SelectedIndexChanged gọi TaoMSSV để gợi ý mã có tiền tố phù hợp."],
    ["Tải lại ở Chủ đề 2 khi chưa lưu", "Danh sách trong bộ nhớ bị thay bằng dữ liệu DSNV.txt; thay đổi chưa ghi bị bỏ."],
    ["Tệp dữ liệu không tồn tại", "Hàm đọc trả về danh sách rỗng; khi trình diễn cần kiểm tra tệp được sao chép cạnh ứng dụng."],
])

doc.add_heading("Kết luận", level=1)
para("Hai ứng dụng đã mở và hiển thị dữ liệu mẫu. Chủ đề 2 trình diễn chọn bản ghi và cập nhật trên giao diện; Chủ đề 3 trình diễn nạp JSON và tìm kiếm theo tên. Mã nguồn được tổ chức để có thể lần theo một thao tác từ control, qua danh sách trong bộ nhớ, đến tệp dữ liệu.")

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run("Báo cáo thực hành giữa kỳ  Chủ đề 2 và Chủ đề 3")

doc.save(OUT)
print(OUT)
